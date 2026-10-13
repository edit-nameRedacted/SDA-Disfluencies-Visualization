import csv, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
src, out = sys.argv[1], sys.argv[2]
rows = list(csv.DictReader(open(src, encoding="utf8"))); N = len(rows)
CODES = ["Repair initiation", "Not repair", "Cant tell"]
base = Font(name="Arial", size=10); bold = Font(name="Arial", size=10, bold=True); head = Font(name="Arial", bold=True, size=11); title = Font(name="Arial", bold=True, size=14)
yellow = PatternFill("solid", start_color="FFF6CC"); grey = PatternFill("solid", start_color="EEEEEA")
wrap = Alignment(wrap_text=True, vertical="top"); thin = Border(bottom=Side(style="thin", color="BBBBBB"))
wb = Workbook(); ws = wb.active; ws.title = "Read me"
ws.column_dimensions["A"].width = 28; ws.column_dimensions["B"].width = 90
r = 1
def put(a, b=None, font=base, bfont=base, fill=None):
    global r
    ws.cell(row=r, column=1, value=a).font = font; ws.cell(row=r, column=2, value=b).font = bfont
    for c in (1, 2):
        ws.cell(row=r, column=c).alignment = wrap
        if fill: ws.cell(row=r, column=c).fill = fill
    r += 1
put("Repair initiations: coding sheet", font=title); r += 1
put("The question", "For each row: in the Line, does the speaker signal a problem hearing or understanding what the other person just said (the Before line), so that the other person is asked to say it again or explain it?", font=bold)
put("What to do", f"On the Coding sheet, read Before, Line and After, then type one code in column E for each of the {N} rows: Repair initiation, Not repair, or Cant tell. Notes (column F) are optional.", font=bold)
put("Stay blind", "Do not open repair_coding_key.csv until you have finished. It shows which rows the script flagged.", font=bold)
r += 1
put("Progress", font=head)
put("Rows coded", f"=COUNTA(Coding!E2:E{N+1})", font=bold)
first = r
for code in CODES: put(code, f"=COUNTIF(Coding!E$2:E${N+1},A{r})", font=bold)
put("Something else typed", f"=B{first-1}-SUM(B{first}:B{r-1})", font=bold)
r += 1
put("Code it Repair initiation when", font=head)
for a, b in [
 ("Open request", "The Line only says something went wrong, without saying what: \"what?\", \"huh?\", \"pardon?\", \"say again?\", \"sorry?\". After usually repeats Before."),
 ("Targeted request", "The Line points at a part of Before: repeats a word with question intonation (\"Tucson?\"), or asks about one part (\"North where?\", \"the what?\")."),
 ("Understanding check", "The Line offers its reading of Before for confirmation: \"you mean the bus, not the train?\". Count it when it checks what was meant, not when it adds a new point."),
 ("Signs in After", "After repeats, rephrases, clarifies or confirms what was said in Before. This supports the code but is not required: sometimes the reply is just \"yeah\".")]:
    put(a, b, font=bold)
r += 1
put("Code it Not repair when", font=head)
for a, b in [
 ("A news receipt", "\"Really?\", \"you did?\", \"oh is it?\", \"is that right?\": these react to news that was understood."),
 ("A new question", "\"What kind do you have?\", \"what's that like?\": asks for more, not for a repeat or explanation of something unclear."),
 ("A backchannel or reaction", "\"Uh-huh\", \"huh\" said as a reaction, \"I see\"."),
 ("An apology", "\"I'm sorry\" in sympathy, not \"sorry?\" meaning \"say that again\".")]:
    put(a, b, font=bold)
r += 1
put("Hard cases", "The examples on this sheet are made up, so that none of them is a row you will code.", font=head)
put("\"What's that?\"", "Repair initiation if the Before used a word or name the speaker did not know or catch (\"We use a mandoline.\" \"What's that?\"). Not repair if it asks to hear about something announced (\"Guess what happened today.\" \"What's that?\").", font=bold)
put("A repeat that just agrees", "\"Expensive.\" said as agreement (After: \"yeah\") is Not repair. A repeat that asks (\"Forty?\" After: \"No, fourteen\") is Repair initiation.", font=bold)
put("No punctuation", "The words have no punctuation. Read the line as it would most likely have been said, using Before and After.", font=bold)
r += 1
put("About the sample", f"{N} response-opening lines from the 650 calls: every line flagged by the open, repeat and check rules, 60 of the lines tagged 'signal non-understanding' by the corpus, 60 short unflagged questions, and 60 other unflagged lines. Shuffled; seed 20261013; drawn by make_repair_sample.py.", font=bold)
wc = wb.create_sheet("Coding")
cols = [("Item", 6), ("Before (other speaker)", 50), ("Line (code this)", 50), ("After (other speaker)", 50), ("Code", 20), ("Notes", 30)]
for i, (n, w) in enumerate(cols):
    c = wc.cell(row=1, column=i + 1, value=n); c.font = head; c.fill = grey; c.alignment = wrap
    wc.column_dimensions["ABCDEF"[i]].width = w
for j, row in enumerate(rows, start=2):
    for i, v in enumerate([int(row["item"]), row["before"], row["line"], row["after"], None, None]):
        c = wc.cell(row=j, column=i + 1, value=v); c.font = bold if i == 2 else base; c.alignment = wrap; c.border = thin
        if i >= 4: c.fill = yellow
wc.freeze_panes = "B2"
dv = DataValidation(type="list", formula1='"' + ",".join(CODES) + '"', allow_blank=True, showErrorMessage=False)
wc.add_data_validation(dv); dv.add(f"E2:E{N+1}")
wb.save(out); print("saved", out, N)
