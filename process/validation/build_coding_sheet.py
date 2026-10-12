import csv, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

src, out = sys.argv[1], sys.argv[2]
rows = list(csv.DictReader(open(src, encoding="utf8")))
N = len(rows)
CODES = ["Complete fluent", "Complete disfluent", "Unfinished", "Fragment", "Response token", "Cant tell"]

wb = Workbook()
head = Font(name="Arial", bold=True, size=11); base = Font(name="Arial", size=10); bold = Font(name="Arial", size=10, bold=True)
title = Font(name="Arial", bold=True, size=14)
fill_in = PatternFill("solid", start_color="FFF6CC"); grey = PatternFill("solid", start_color="EEEEEA")
wrap = Alignment(wrap_text=True, vertical="top")
thin = Border(bottom=Side(style="thin", color="BBBBBB"))

# --- Read me ------------------------------------------------------------------------------
ws = wb.active; ws.title = "Read me"
ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 70; ws.column_dimensions["C"].width = 44
r = 1
def put(a, b=None, c=None, font=base, fill=None):
    global r
    for col, v in zip("ABC", (a, b, c)):
        cell = ws[f"{col}{r}"]; cell.value = v; cell.font = font; cell.alignment = wrap
        if fill: cell.fill = fill
    r += 1
put("Sentence types: coding sheet, second sample", font=title); r += 1
put("What to do", f"Go to the Coding sheet. For each of the {N} rows, read the Utterance column and type or pick one code in the Code column. The two lines before and the line after are context only.", font=base)
ws[f"A{r-1}"].font = bold
put("The six codes", "Type them exactly as written below, or pick them from the dropdown if your app shows one. Letters only, no commas.", font=base); ws[f"A{r-1}"].font = bold
put("Cells to fill in", "The yellow cells on the Coding sheet: Code (column F) and, if you wish, Notes (column G).", font=base); ws[f"A{r-1}"].font = bold
put("Stay blind", "Do not open coding_key_2.csv, and do not look the utterances up in the piece, until you have finished. Either shows the script's answer.", font=base); ws[f"A{r-1}"].font = bold
put("Full codebook", "The Sentence-Type Codebook document has the definitions, the decision steps and the hard cases. The short version is below.", font=base); ws[f"A{r-1}"].font = bold
r += 1
put("Progress", font=head)
put("Rows coded", f"=COUNTA(Coding!F2:F{N+1})", f"of {N}")
first = r
for code in CODES:
    put(code, f'=COUNTIF(Coding!F$2:F${N+1},A{r})')
put("Something else typed", f"=B{first-1}-SUM(B{first}:B{r-1})", "should be 0; check the spelling of these rows")
r += 1
put("The codes", font=head)
put("Code", "What it is", "Examples", font=bold, fill=grey)
for a, b, c in [
    ("Complete fluent", "A whole sentence: a subject with its verb, finished, with no disfluency. Questions and commands count. An opening and, but, so or because does not stop it being whole.", "Well what are your hobbies / and they left the dog in the car / Because I look at him"),
    ("Complete disfluent", "A whole sentence, finished, with at least one disfluency: uh or um, a repeated or reworded stretch, a cut-off word, you know, I mean.", "I uh was just in a supermarket / you know they weren't a state / I mean the system isn't perfect"),
    ("Unfinished", "The speaker started something and broke it off or let it trail away.", "and then you um / Um our our friend has a / how wha-"),
    ("Fragment", "No main clause: no verb at all, a verb with its subject left out, a clause opening with if, when or that standing alone, or nothing but linking words.", "In New Mexico / A hamster / if it's not raining / Got a bunch of groves in it / and uh"),
    ("Response token", "Nothing but a reaction or answer word. A lone uh or um counts here.", "Yeah / Uh-huh / Right / Oh yeah"),
    ("Cant tell", "The words do not let you decide. Say why in Notes.", None)]:
    put(a, b, c); ws[f"A{r-1}"].font = bold
r += 1
put("Decision steps", font=head)
for a, b in [
    ("Step 1", "Nothing but a reaction or answer word (yeah, uh-huh, right, okay, oh, no, a lone uh or um)? Response token."),
    ("Step 2", "Nothing but linking words (and, but, so, or, then, because), with or without uh or um? Fragment."),
    ("Step 3", "Did the speaker break off or trail away before finishing what they started? Unfinished."),
    ("Step 4", "No main clause (no verb; or a verb with no subject; or only an if, when or that clause)? Fragment. Ignore an opening and, but, so or because when you judge this."),
    ("Step 5", "A disfluency anywhere (uh, um, a repeat or rewording, a cut-off word, you know, I mean)? Complete disfluent. Otherwise Complete fluent."),
    ("More than one sentence", "Judge the longest one."),
    ("Not disfluencies", "Opening words well, so, and, but, oh; like, kind of, I guess; you know used as part of the sentence (you know the one); a word repeated for emphasis (very very good).")]:
    put(a, b); ws[f"A{r-1}"].font = bold
r += 1
put("About the sample", f"{N} utterances, 60 from each type the script assigned under the revised rules, drawn at random from the 650 calls with a fixed seed (20261011) by make_coding_sample.py, then shuffled. None of them was in the first sample. Words are shown without punctuation or markup. Utterances the other speaker cut into are joined into one line.")
ws[f"A{r-1}"].font = bold

# --- Coding --------------------------------------------------------------------------------
wc = wb.create_sheet("Coding")
cols = [("Item", 6), ("Two before", 36), ("One before", 36), ("Utterance (code this line)", 60), ("One after", 36), ("Code", 20), ("Notes", 30)]
for i, (name, width) in enumerate(cols):
    c = wc.cell(row=1, column=i + 1, value=name); c.font = head; c.fill = grey; c.alignment = wrap
    wc.column_dimensions["ABCDEFG"[i]].width = width
for j, row in enumerate(rows, start=2):
    vals = [int(row["item"]), row["before_2"], row["before_1"], row["utterance"], row["after_1"], None, None]
    for i, v in enumerate(vals):
        c = wc.cell(row=j, column=i + 1, value=v); c.font = bold if i == 3 else base; c.alignment = wrap; c.border = thin
        if i >= 5: c.fill = fill_in
wc.freeze_panes = "B2"
dv = DataValidation(type="list", formula1='"' + ",".join(CODES) + '"', allow_blank=True, showErrorMessage=False)
wc.add_data_validation(dv); dv.add(f"F2:F{N+1}")
wb.save(out)
print("saved", out, N)
